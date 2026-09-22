"""Shipping-mark workbook generation.

Reads a supplier packing list, maps each row to a labelled form using the
selected vendor *profile*, fetches barcode images concurrently, and writes a
formatted ``.xlsx`` where rows are grouped into worksheets by the profile's
``sheet_col``.
"""

from __future__ import annotations

import concurrent.futures
import os
import re
import time
from typing import Callable, Optional

import openpyxl
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import column_index_from_string

from . import config
from .barcode import BarcodeClient

Log = Callable[[str], None]

_TEMPLATE_RE = re.compile(r"\{([A-Za-z]+)\}")
_COLUMN_RE = re.compile(r"[A-Za-z]+")
_INVALID_SHEET_CHARS = re.compile(r"[\[\]:\*\?\/\\]")


# --------------------------------------------------------------------------- #
# Field-mapping helpers
# --------------------------------------------------------------------------- #
def column_value(row: tuple, letter: str) -> object:
    """Return the value of a column letter for a ``values_only`` row tuple."""
    try:
        index = column_index_from_string(str(letter).strip()) - 1
    except (ValueError, AttributeError):
        return None
    return row[index] if 0 <= index < len(row) else None


def render_template(template: str, row: tuple) -> str:
    """Replace every ``{COL}`` token with the column's value (None -> '')."""
    def repl(match: re.Match) -> str:
        value = column_value(row, match.group(1))
        return "" if value is None else str(value)

    return _TEMPLATE_RE.sub(repl, str(template))


def resolve_field(spec: str, row: tuple) -> str:
    """Resolve any field spec against a row, supporting all four input formats:

    1. Single column   ->  ``B``
    2. Merged columns    ->  ``I+J``          (non-empty values, space separated)
    3. Template          ->  ``{A} size {B}`` (substitute ``{COL}``, keep text)
    4. Literal text      ->  ``1PCS OF 1 CTN``

    The format is detected from the spec itself, so every field accepts all
    four without the caller needing to know which one was used.
    """
    spec = "" if spec is None else str(spec)
    if not spec:
        return ""
    # Template: any {COL} token -> substitute and keep the surrounding text.
    if _TEMPLATE_RE.search(spec):
        return render_template(spec, row)
    # Single / merged columns: every '+'-separated part is a column letter.
    parts = [p.strip() for p in spec.split("+")]
    if parts and all(_COLUMN_RE.fullmatch(p) for p in parts):
        values = [
            str(value).strip()
            for value in (column_value(row, part) for part in parts)
            if value is not None and str(value).strip()
        ]
        return " ".join(values)
    # Literal text (anything else) is used verbatim.
    return spec


def unique_sheet_name(name: str, used: set[str]) -> str:
    """Sanitise to Excel's rules (<=31 chars, no []:*?/\\, unique)."""
    name = _INVALID_SHEET_CHARS.sub("_", str(name)).strip() or "Sheet"
    name = name[:31]
    base, counter = name, 1
    while name.lower() in used:
        suffix = f"_{counter}"
        name = base[: 31 - len(suffix)] + suffix
        counter += 1
    used.add(name.lower())
    return name


# --------------------------------------------------------------------------- #
# Generator
# --------------------------------------------------------------------------- #
class ShippingMarkGenerator:
    def __init__(self, log: Log = print) -> None:
        self._log = log
        self._barcodes = BarcodeClient(log=log)
        self._border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin"),
        )
        self._font = Font(name=config.FONT_NAME, size=config.FONT_SIZE, bold=config.FONT_BOLD)

    # -- public API -------------------------------------------------------- #
    def generate(self, profile: dict, source_path: str, po_no: str, output_path: str) -> str:
        started = time.time()
        self._log(f"Reading source workbook: {source_path}")
        source = openpyxl.load_workbook(source_path, data_only=True).active

        rows = list(enumerate(source.iter_rows(min_row=2, values_only=True), start=2))
        self._prefetch_barcodes(rows, profile)

        self._log("Building worksheets...")
        workbook = Workbook()
        workbook.remove(workbook.active)
        self._write_rows(workbook, rows, profile, po_no)

        if not workbook.worksheets:
            workbook.create_sheet(title="Sheet1")

        if os.path.exists(output_path):
            os.remove(output_path)
        workbook.save(output_path)

        elapsed = time.time() - started
        forms = len(rows)
        images = sum(1 for _, r in rows if resolve_field(profile["item_code_col"], r))
        self._log("")
        self._log(f"Done in {elapsed:.1f}s")
        self._log(f"Generated {forms} forms and {images} barcode lookups")
        self._log(f"Saved: {output_path}")
        return output_path

    # -- internals --------------------------------------------------------- #
    def _prefetch_barcodes(self, rows, profile) -> None:
        skus = [resolve_field(profile["item_code_col"], r) for _, r in rows]
        skus = [s for s in skus if s]
        self._log(f"Fetching {len(skus)} barcode images ({config.BARCODE_MAX_WORKERS} workers)...")
        with concurrent.futures.ThreadPoolExecutor(max_workers=config.BARCODE_MAX_WORKERS) as pool:
            list(pool.map(self._barcodes.get, skus))

    def _write_rows(self, workbook, rows, profile, po_no) -> None:
        last_color: Optional[str] = None
        last_split: Optional[str] = None
        sheet = None
        row_cursor = config.FORM_START_ROW
        used_names: set[str] = set()

        for index, row in rows:
            split_raw = resolve_field(profile["sheet_col"], row)
            split_key = split_raw or None

            # New worksheet only when the split value *changes*; blank or
            # repeated values keep stacking on the current sheet.
            if sheet is None:
                name = unique_sheet_name(split_key or f"Sheet_{index}", used_names)
                sheet = self._new_sheet(workbook, name)
                row_cursor, last_split = config.FORM_START_ROW, split_key
            elif split_key is not None and split_key != last_split:
                sheet = self._new_sheet(workbook, unique_sheet_name(split_key, used_names))
                row_cursor, last_split = config.FORM_START_ROW, split_key

            color = resolve_field(profile["color_col"], row)
            if color:
                last_color = color
            else:
                color = last_color or ""

            item_code = resolve_field(profile["item_code_col"], row)
            form = self._build_form(profile, row, po_no, item_code, color)

            form_top = row_cursor
            row_cursor = self._write_form(sheet, form, row_cursor)
            row_cursor = self._write_made_in(sheet, row_cursor)

            if item_code:
                self._add_barcode(sheet, item_code, form_top, len(form))

            for _ in range(config.EMPTY_ROWS_AFTER_FORM):
                sheet.row_dimensions[row_cursor].height = config.ROW_HEIGHT
                row_cursor += 1

    @staticmethod
    def _build_form(profile, row, po_no, item_code, color) -> list[tuple[str, object]]:
        desc = resolve_field(profile["description_col"], row)
        gross = resolve_field(profile["gross_wt_col"], row)
        net = resolve_field(profile["net_wt_col"], row)
        return [
            ("PO NO.:", po_no),
            ("ITEM CODE:", item_code),
            ("DESCRIPTION:", desc.upper()),
            ("PACK SIZE:", resolve_field(profile["pack_size_tpl"], row)),
            ("COLOR:", color.upper() if color else ""),
            ("CTN PER PCS:", resolve_field(profile["ctn_per_pcs"], row)),
            ("CTN NO.:", resolve_field(profile["ctn_no_tpl"], row)),
            ("GROSS WT:", f"{gross} KG" if gross else ""),
            ("NET WT:", f"{net} KG" if net else ""),
        ]

    def _new_sheet(self, workbook, name: str):
        sheet = workbook.create_sheet(title=name)
        sheet.column_dimensions["A"].width = config.COL_A_WIDTH
        sheet.column_dimensions["B"].width = config.COL_B_WIDTH
        return sheet

    def _write_form(self, sheet, form, start_row: int) -> int:
        row = start_row
        for label, value in form:
            label_cell = sheet.cell(row=row, column=1, value=label)
            value_cell = sheet.cell(row=row, column=2, value=value)
            for cell in (label_cell, value_cell):
                cell.border = self._border
                cell.alignment = Alignment(horizontal="left")
                cell.font = self._font
            sheet.row_dimensions[row].height = config.ROW_HEIGHT
            row += 1
        return row

    def _write_made_in(self, sheet, row: int) -> int:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        cell = sheet.cell(row=row, column=1, value="MADE IN PRC")
        cell.alignment = Alignment(horizontal="left")
        cell.font = self._font
        for col in (1, 2):
            sheet.cell(row=row, column=col).border = self._border
        sheet.row_dimensions[row].height = config.ROW_HEIGHT
        return row + 1

    def _add_barcode(self, sheet, sku, top_row: int, field_count: int) -> None:
        image = self._barcodes.get(sku)
        if image is None:
            self._log(f"[{sku}] no barcode image")
            return
        try:
            height_px = int((field_count + 1) * config.ROW_HEIGHT * config.IMAGE_HEIGHT_MULTIPLIER)
            xl_image = self._to_xl_image(image)
            xl_image.height = height_px
            xl_image.width = int(image.width * (height_px / image.height))
            sheet.add_image(xl_image, f"{config.IMAGE_COLUMN}{top_row}")
            self._log(f"[{sku}] barcode placed on {sheet.title}")
        except Exception as exc:  # noqa: BLE001
            self._log(f"[{sku}] could not place barcode: {exc}")

    @staticmethod
    def _to_xl_image(pil_image) -> XLImage:
        import io as _io
        buffer = _io.BytesIO()
        pil_image.save(buffer, format="PNG")
        buffer.seek(0)
        return XLImage(buffer)
